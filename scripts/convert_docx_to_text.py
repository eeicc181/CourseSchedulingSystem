#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
将指定的Word文档转换为纯文本格式
"""

from pathlib import Path
from docx import Document

# 需要转换的文件列表
FILES_TO_CONVERT = [
    r"D:\AAA_CodeLearning\CourseArrange\docs\孙传家+许东徽+柏璐瑶+梁治国+排课系统\孙传家+许东徽+柏璐瑶+梁治国+排课系统\01-《软件工程专业专业生产实习》任务书.docx",
    r"D:\AAA_CodeLearning\CourseArrange\docs\孙传家+许东徽+柏璐瑶+梁治国+排课系统\孙传家+许东徽+柏璐瑶+梁治国+排课系统\02-第20组《基于遗传算法的自动排课系统》报告书.docx",
    r"D:\AAA_CodeLearning\CourseArrange\docs\孙传家+许东徽+柏璐瑶+梁治国+排课系统\孙传家+许东徽+柏璐瑶+梁治国+排课系统\03-团队交流记录表.docx",
    r"D:\AAA_CodeLearning\CourseArrange\docs\孙传家+许东徽+柏璐瑶+梁治国+排课系统\孙传家+许东徽+柏璐瑶+梁治国+排课系统\04-《软件工程专业专业生产实习》指导过程记录表.docx",
    r"D:\AAA_CodeLearning\CourseArrange\docs\孙传家+许东徽+柏璐瑶+梁治国+排课系统\孙传家+许东徽+柏璐瑶+梁治国+排课系统\05-《软件工程专业生产实习》答辩记录表.doc",
]

# 输出目录
OUTPUT_DIR = Path(r"D:\AAA_CodeLearning\CourseArrange\docs")


def convert_docx_to_text(docx_path, output_dir):
    """将Word文档转换为纯文本"""
    try:
        docx_path = Path(docx_path)
        
        # 检查文件是否存在
        if not docx_path.exists():
            return False, f"文件不存在: {docx_path}"
        
        # 读取Word文档
        doc = Document(docx_path)
        
        # 提取所有段落文本
        text_content = []
        for paragraph in doc.paragraphs:
            text_content.append(paragraph.text)
        
        # 提取表格内容
        for table in doc.tables:
            for row in table.rows:
                row_text = []
                for cell in row.cells:
                    row_text.append(cell.text.strip())
                text_content.append(" | ".join(row_text))
        
        # 生成输出文件名
        output_file = output_dir / (docx_path.stem + ".txt")
        
        # 写入文本文件
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write("=" * 80 + "\n")
            f.write(f"源文件: {docx_path.name}\n")
            f.write("=" * 80 + "\n\n")
            f.write("\n".join(text_content))
        
        return True, docx_path.name
    
    except Exception as e:
        return False, f"{docx_path.name}: {str(e)}"


def main():
    """主函数"""
    print("=" * 80)
    print("开始转换Word文档到纯文本格式")
    print("=" * 80)
    print(f"输出目录: {OUTPUT_DIR}\n")
    
    success_count = 0
    failed_files = []
    
    for i, file_path in enumerate(FILES_TO_CONVERT, 1):
        success, info = convert_docx_to_text(file_path, OUTPUT_DIR)
        
        if success:
            success_count += 1
            print(f"[{i}/{len(FILES_TO_CONVERT)}] ✓ {info}")
        else:
            failed_files.append(info)
            print(f"[{i}/{len(FILES_TO_CONVERT)}] ✗ {info}")
    
    # 输出统计信息
    print("\n" + "=" * 80)
    print("转换完成!")
    print("=" * 80)
    print(f"成功: {success_count} 个文件")
    print(f"失败: {len(failed_files)} 个文件")
    print(f"输出目录: {OUTPUT_DIR}")
    
    if failed_files:
        print("\n失败的文件:")
        for failed in failed_files:
            print(f"  - {failed}")


if __name__ == "__main__":
    main()
